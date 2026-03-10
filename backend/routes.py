"""
FastAPI routes for the research agent API.
Config, pipeline SSE, notebook, export, save/load, agent re-run.
"""

import json
import io
import os
import tempfile
from datetime import datetime
from fastapi import APIRouter, HTTPException
from fastapi.responses import StreamingResponse, Response
from pydantic import BaseModel, Field
from typing import Optional

from backend.config import config
from backend.pipeline import run_pipeline, rerun_agent
from backend.tools.code_executor import execute_code
from backend.tools.llm_client import call_llm_chat
from backend.tools import user_choice
from backend.models import ChatRequest

router = APIRouter(prefix="/api")


# ── Config ────────────────────────────────────────────────────

class ConfigUpdate(BaseModel):
    openai_api_key: Optional[str] = None
    openai_base_url: Optional[str] = None
    openai_model: Optional[str] = None
    max_papers: Optional[int] = None

class ConfigStatus(BaseModel):
    is_configured: bool
    openai_base_url: str
    openai_model: str
    max_papers: int
    has_api_key: bool

@router.post("/config")
async def update_config(update: ConfigUpdate):
    config.update(**update.model_dump(exclude_none=True))
    return {"status": "ok", "is_configured": config.is_configured()}

@router.get("/config/status", response_model=ConfigStatus)
async def get_config_status():
    return ConfigStatus(
        is_configured=config.is_configured(),
        openai_base_url=config.openai_base_url,
        openai_model=config.openai_model,
        max_papers=config.max_papers,
        has_api_key=bool(config.openai_api_key),
    )


# ── Research Pipeline ─────────────────────────────────────────

@router.get("/research/start")
async def start_research(topic: str):
    async def precheck_stream():
        from backend.models import PipelineEvent
        if not topic or len(topic.strip()) < 3:
            evt = PipelineEvent(stage="done", status="error", message="Topic too short.")
            yield f"data: {evt.model_dump_json()}\n\n"
            return
        if not config.is_configured():
            evt = PipelineEvent(stage="done", status="error", message="API keys not configured. Please open settings.")
            yield f"data: {evt.model_dump_json()}\n\n"
            return
        
        # Valid, proceed to real pipeline
        import backend.pipeline
        async for event in backend.pipeline.run_pipeline(topic.strip()):
            yield event

    return StreamingResponse(
        precheck_stream(),
        media_type="text/event-stream",
        headers={"Cache-Control": "no-cache", "Connection": "keep-alive", "X-Accel-Buffering": "no"},
    )


# ── Notebook Execution ────────────────────────────────────────

class ExecuteRequest(BaseModel):
    code: str
    timeout: int = 300

@router.post("/notebook/execute")
async def execute_cell(req: ExecuteRequest):
    result = await execute_code(req.code, timeout=min(req.timeout, 300))
    return {
        "success": result.success, "stdout": result.stdout, "stderr": result.stderr,
        "images": result.images, "error": result.error, "execution_time": result.execution_time,
    }


# ── Agent Chat ────────────────────────────────────────────────

@router.post("/notebook/chat")
async def agent_chat(req: ChatRequest):
    if not config.is_configured():
        raise HTTPException(status_code=400, detail="API key not configured.")

    agent_prompts = {
        "planner": """You are the Planner Agent. Help the user refine research plans, suggest new angles,
adjust methodology, or explain plan scores. If asked to modify plans, describe the changes clearly.
You can suggest re-running with specific feedback.""",
        "paper_reader": """You are the Paper Reader Agent. Help the user understand paper findings,
suggest additional papers to read, compare methodologies across papers, and discuss the literature synthesis.
If asked for more papers, suggest specific search queries.""",
        "hypothesis_gen": """You are the Hypothesis Agent. Help the user refine hypotheses, suggest improvements,
explain the rationale behind each hypothesis, discuss feasibility. If asked to improve a hypothesis,
provide specific technical modifications.""",
        "experiment": """You are the Experiment Agent in a Jupyter notebook. Help the user:
- Debug code errors, improve models, explain results
- If asked for code changes, provide COMPLETE corrected code in code blocks
- Suggest data augmentation, hyperparameter tuning, model architecture changes""",
        "evaluation": """You are the Evaluation Agent. Help the user understand metrics, suggest additional
evaluation approaches, discuss model strengths/weaknesses, and propose improvements to methodology.""",
        "writer": """You are the Writer Agent for research papers. Help the user improve report sections,
rephrase content, add citations, expand methodology descriptions, or restructure the paper.""",
    }

    system_msg = agent_prompts.get(req.agent,
        "You are an AI Research Assistant. Help the user with their research tasks. Be concise and actionable.")

    messages = [{"role": "system", "content": system_msg}]
    if req.context:
        messages.append({"role": "system", "content": f"Current context:\n{req.context[:6000]}"})
    messages.append({"role": "user", "content": req.message})
    try:
        response = await call_llm_chat(messages, max_tokens=3000)
        return {"response": response}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


# ── Agent Re-run ──────────────────────────────────────────────

class RerunRequest(BaseModel):
    stage: str
    feedback: str = ""
    current_state: dict = Field(default_factory=dict)

@router.post("/agent/rerun")
async def agent_rerun(req: RerunRequest):
    """Re-run a specific agent with user feedback."""
    if not config.is_configured():
        raise HTTPException(status_code=400, detail="API key not configured.")
    try:
        result = await rerun_agent(req.stage, req.feedback, req.current_state)
        return {"status": "ok", "stage": req.stage, "data": result}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


# ── User Choice (Plan/Hypothesis selection) ───────────────────

class ChoiceRequest(BaseModel):
    choice_id: str
    selected_index: int

@router.post("/choice")
async def submit_choice(req: ChoiceRequest):
    """Submit user's choice for plan or hypothesis selection."""
    success = user_choice.submit_choice(req.choice_id, req.selected_index)
    if not success:
        raise HTTPException(status_code=404, detail="No pending choice with that ID.")
    return {"status": "ok", "choice_id": req.choice_id, "selected_index": req.selected_index}


# ── Export: Jupyter Notebook (.ipynb) ─────────────────────────

class NotebookExportRequest(BaseModel):
    cells: list[dict] = Field(default_factory=list)
    topic: str = "Research Experiment"

@router.post("/export/ipynb")
async def export_ipynb(req: NotebookExportRequest):
    """Convert notebook cells to Jupyter .ipynb format."""
    nb_cells = []
    for cell in req.cells:
        nb_cells.append({
            "cell_type": "markdown",
            "metadata": {},
            "source": [f"## {cell.get('title', 'Cell')}\n", cell.get('description', '')]
        })
        nb_cells.append({
            "cell_type": "code",
            "metadata": {},
            "source": cell.get("code", "").split("\n"),
            "outputs": [],
            "execution_count": None,
        })

    notebook = {
        "nbformat": 4, "nbformat_minor": 5,
        "metadata": {
            "kernelspec": {"display_name": "Python 3", "language": "python", "name": "python3"},
            "language_info": {"name": "python", "version": "3.10.0"},
        },
        "cells": nb_cells,
    }

    content = json.dumps(notebook, indent=2)
    return Response(
        content=content,
        media_type="application/json",
        headers={"Content-Disposition": f'attachment; filename="experiment_{datetime.now().strftime("%Y%m%d_%H%M%S")}.ipynb"'},
    )


# ── Export: Word Document (.docx) ─────────────────────────────

class DocxExportRequest(BaseModel):
    title: str = "Research Report"
    sections: dict = Field(default_factory=dict)

@router.post("/export/docx")
async def export_docx(req: DocxExportRequest):
    """Export report as .docx file."""
    try:
        from docx import Document
        from docx.shared import Pt, Inches
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError:
        raise HTTPException(status_code=500, detail="python-docx not installed. Run: pip install python-docx")

    doc = Document()

    # Title
    title_para = doc.add_heading(req.title, level=0)
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Sections
    section_order = ['abstract', 'introduction', 'literature_review', 'methodology',
                     'experiment_setup', 'results', 'discussion', 'conclusion', 'references']

    for key in section_order:
        content = req.sections.get(key, "")
        if content:
            heading = key.replace("_", " ").title()
            doc.add_heading(heading, level=1)
            for para_text in content.split("\n\n"):
                if para_text.strip():
                    p = doc.add_paragraph(para_text.strip())
                    p.style.font.size = Pt(11)

    # Save to bytes
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="report_{datetime.now().strftime("%Y%m%d_%H%M%S")}.docx"'},
    )


# ── Save/Load Progress ───────────────────────────────────────

@router.post("/progress/save")
async def save_progress(data: dict):
    """Save returns the data back for client-side download."""
    return data

@router.get("/health")
async def health_check():
    return {"status": "healthy", "configured": config.is_configured()}
